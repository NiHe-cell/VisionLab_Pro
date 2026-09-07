# -*- coding: utf-8 -*-
"""Sole VisionLab Pro task-summary report helpers.

Canonical file (do not create another .docx):
  %USERPROFILE%/Desktop/VisionLab_Pro任务总结报告.docx

After every completed Px-Tyy task, append to that file with these helpers.
Never write a backup copy or a separate markdown summary.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

REPORT_NAME = "VisionLab_Pro任务总结报告.docx"
YAHEI = "微软雅黑"
CONSOLAS = "Consolas"


def locate_report() -> Path:
    path = Path.home() / "Desktop" / REPORT_NAME
    if not path.is_file():
        raise FileNotFoundError(
            f"Canonical task report not found: {path}. "
            "Do not create a new document; restore this file first."
        )
    return path


def open_report() -> Document:
    return Document(str(locate_report()))


def save_report(doc: Document) -> Path:
    path = locate_report()
    doc.save(str(path))
    return path


def set_run_font(run, name=YAHEI, size=None, bold=None, east_asia=YAHEI):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    if east_asia:
        rFonts.set(qn("w:eastAsia"), east_asia)


def add_styled(doc, text, style="Normal", bold=None, font=YAHEI, size=None, space_after=None):
    p = doc.add_paragraph()
    p.style = doc.styles[style]
    if space_after is not None:
        p.paragraph_format.space_after = Emu(space_after)
    run = p.add_run(text)
    set_run_font(run, name=font, size=size, bold=bold, east_asia=YAHEI if font == YAHEI else None)
    if font != YAHEI:
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), YAHEI)
    return p


def h1(doc, text):
    add_styled(doc, text, style="Heading 1")


def h2(doc, text):
    add_styled(doc, text, style="Heading 2")


def body(doc, text):
    add_styled(doc, text, style="Normal", bold=False)


def label(doc, text):
    add_styled(doc, text, style="Normal", bold=True)


def code(doc, text):
    add_styled(doc, text, style="Normal", font=CONSOLAS, size=9, space_after=76200)


def bullet(doc, text):
    add_styled(doc, text, style="List Bullet")


def add_grid_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.autofit = True
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(header)
        set_run_font(run, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, value in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(value)
            set_run_font(run, bold=False)
    return table


def clear_runs(paragraph):
    for child in list(paragraph._element):
        if child.tag == qn("w:r"):
            paragraph._element.remove(child)


def replace_paragraph_text(paragraph, lines, bold=True, size=None):
    clear_runs(paragraph)
    for i, line in enumerate(lines):
        if i:
            paragraph.add_run("\n")
        run = paragraph.add_run(line)
        set_run_font(run, bold=bold, size=size)


def update_cover(doc, phase_range: str, theme_line: str):
    """phase_range like '1–5'; theme_line is the second title line."""
    replace_paragraph_text(
        doc.paragraphs[0],
        [
            f"VisionLab Pro — Phase {phase_range} 任务总结报告",
            theme_line,
        ],
        bold=True,
        size=22,
    )
